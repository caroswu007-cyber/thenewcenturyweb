# -*- coding: utf-8 -*-
"""One-shot: write 中文 column for 03 / 06 workbooks from curated maps (see historic page-copy fill)."""
from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
PAGE_COPY = ROOT / "docs" / "page-copy"

FOUNDER_ZH: dict[str, str] = {
    "founderStorySurfaceCopy.heroNamesLine": "John Long Woo · Caros Woo · Sam",
    "founderStorySurfaceCopy.backToAbout": "← 组织概览",
    "founderStorySurfaceCopy.legacyTimelineLink": "旧站上的早期分阶段时间线",
    "founderStoryPage.heroTitle": "我们的故事纲要 — 2026",
    "founderStoryPage.heroBadge": "创始人叙事",
    "founderStoryPage.intro": (
        "我是 **John Long Woo**。与我的两个儿子 **Caros Woo**、**Sam**——我们 **吴氏父子** 三人——"
        "所经历的涉及 **灵虚世界** 的事件已多得无法计数。这些经历使我们成为 **通往灵虚世界之门的首要开启者**，"
        "以及 **灵体生命真相的首要探索者**。"
    ),
    "founderStoryPage.truths.title": "三个真相",
    "founderStoryPage.truths.items[0].label": "真相一",
    "founderStoryPage.truths.items[0].blocks[0].title": "肉体与灵体",
    "founderStoryPage.truths.items[0].blocks[0].text": (
        "活人由 **肉体** 与 **灵体**（spirit / soul）共同构成。**鬼魂**（spirit / Ghost）在人死后仍然存在。"
    ),
    "founderStoryPage.truths.items[0].blocks[1].title": "寄生与疾病",
    "founderStoryPage.truths.items[0].blocks[1].text": (
        "对数以亿计的它们而言，生存的唯一方式是 **附体于人** 并汲取能量。这一模式与 **癫痫**、**抑郁**、"
        "**精神分裂症**、**躯体症状障碍** 等广泛疾病交织在一起。"
    ),
    "founderStoryPage.truths.items[0].blocks[2].title": "灵魂医学",
    "founderStoryPage.truths.items[0].blocks[2].text": (
        "通过 **高级控制灵** 对 **鬼魂**（spirit / Ghost）加以管理，上述状况往往能迅速改善。"
        "基于这项工作，我们创立了 **灵魂医学**。"
    ),
    "founderStoryPage.truths.items[0].blocks[3].title": "档案",
    "founderStoryPage.truths.items[0].blocks[3].text": (
        "更多文献见 **吴氏灵体档案**（Woos Record of Soul）与 **吴氏灵魂医学**（Woos Spirit Medicine）视频档案。"
    ),
    "founderStoryPage.truths.items[1].label": "真相二",
    "founderStoryPage.truths.items[1].blocks[0].text": (
        "**灵虚世界** 有其治理者——强大的 **高级控制灵**；它们由普通 **灵体** 经 **万有元神母体** 升级而来。"
    ),
    "founderStoryPage.truths.items[2].label": "真相三",
    "founderStoryPage.truths.items[2].blocks[0].text": (
        "**万有元神母体**（Universal Matrix of Meta Awareness）是真正的 **创造者**；"
        "**我们的宇宙是其躯体的一部分**。"
    ),
    "founderStoryPage.truths.items[2].blocks[1].title": "档案",
    "founderStoryPage.truths.items[2].blocks[1].text": (
        "更多文献见 **万有元神母体**（Universal Matrix of Meta Awareness）视频档案。"
    ),
    "founderStoryPage.phasesOverview.title": "两大篇章",
    "founderStoryPage.phasesOverview.a.title": "A 阶段 — 2014年5月至2025年2月",
    "founderStoryPage.phasesOverview.a.blocks[0].title": "涉事者",
    "founderStoryPage.phasesOverview.a.blocks[0].text": (
        "我们与附身于我们体内、等级各异的 **附体之灵** 抗争——包括 **C1** 级的 **鬼魂**「**惠安**」层级相关实体、"
        "**C2** 级灵体 **青小仙**，以及 **C3** 级灵体 **钱凯**。它们在我们全家引发了 **严重疾病**。"
    ),
    "founderStoryPage.phasesOverview.a.blocks[1].title": "动机与手段",
    "founderStoryPage.phasesOverview.a.blocks[1].text": (
        "我们唯一的动机是 **自救**：竭尽所能伤害或消灭它们。在此过程中，我们了解了 **鬼魂** 躯体的生理与物理特性，"
        "并通过 **癫痫放电**、**幻听声响**、**意识** 对话、**意识** 操控及相关模式与之沟通——"
        "构筑起一幅经过验证的、宏大的 **灵虚世界** 图景。"
    ),
    "founderStoryPage.phasesOverview.a.blocks[2].title": "遗憾",
    "founderStoryPage.phasesOverview.a.blocks[2].text": (
        "在 **A 阶段** 我们最大的遗憾，是没有系统性地在家人之外的 **其他国家**、**其他人类宿主** 身上接触附体之灵。"
    ),
    "founderStoryPage.phasesOverview.a.blocks[3].title": "记录",
    "founderStoryPage.phasesOverview.a.blocks[3].text": "这一时期我们拍摄了 **吴氏灵体档案**（Woos Record of Soul）。",
    "founderStoryPage.phaseAStages[0].label": "A 1",
    "founderStoryPage.phaseAStages[0].stageTitle": "第 1 阶段 · 欺瞒期 1",
    "founderStoryPage.phaseAStages[0].range": "2018年3月 — 2020年7月",
    "founderStoryPage.phaseAStages[0].paragraphs[0]": (
        "这一时期，我体内的 **低级鬼魂** 利用 **高压静电放电** 作用于饮料与酒，以改变其味道，"
        "旨在伪造一串「科学」现象与法则，最终「证明」**佛教**。"
    ),
    "founderStoryPage.phaseAStages[0].paragraphs[1]": "细节过于繁复，难以用文字完全铺陈。",
    "founderStoryPage.phaseAStages[0].storylineClip.figureAlt": "A1 阶段配图场景",
    "founderStoryPage.phaseAStages[1].label": "A 2",
    "founderStoryPage.phaseAStages[1].stageTitle": "欺瞒期 2",
    "founderStoryPage.phaseAStages[1].range": "2020年7月 — 2020年11月15日",
    "founderStoryPage.phaseAStages[1].paragraphs[0]": (
        "这一时期，**低级鬼魂** 以 **类癫痫发作** 依次假冒 **佛教** 中 **三位最伟大的神（deity）**，"
        "佯称传达 **神意**——直至我与儿子们皈依。"
    ),
    "founderStoryPage.phaseAStages[1].paragraphs[1]": (
        "那是 **荒诞的玩笑**；我们是 **玩物**。即便如此，在这场欺骗中我们仍交换了大量 **哲学洞见**。"
    ),
    "founderStoryPage.phaseAStages[1].storylineClip.figureAlt": "A2 阶段配图场景",
    "founderStoryPage.phaseAStages[2].label": "A 3",
    "founderStoryPage.phaseAStages[2].stageTitle": "驱灵期 1",
    "founderStoryPage.phaseAStages[2].range": "2020年11月15日 — 2021年7月",
    "founderStoryPage.phaseAStages[2].paragraphs[0]": (
        "**2020年11月15日**，我们体内 **鬼魂** 的 **真实身份** 与 **把戏** 终于被揭开。"
        "在恐惧与愤怒中，我们试图 **驱灵**。"
    ),
    "founderStoryPage.phaseAStages[2].paragraphs[1]": (
        "传统方法屡次失败；我们想起它们一贯 **禁止** 我们触碰 **强磁铁**。于是我们开始用磁铁 **伤害** 它们。"
    ),
    "founderStoryPage.phaseAStages[2].paragraphs[2]": (
        "数日内，**长子** 体内的 **鬼魂** 不堪忍受而 **率先离开**；约 **六周** 后，**次子** 体内的 **鬼魂** 亦因同样原因离去。"
    ),
    "founderStoryPage.phaseAStages[2].paragraphs[3]": (
        "**我** 体内的 **鬼魂** 即便受伤仍拒绝离开。我发现 **强静电场** 能大幅削弱其放电：**癫痫** 在那一刻消失——"
        "但 **鬼魂** 仍 **留驻体内**，仿佛在捉迷藏。"
    ),
    "founderStoryPage.phaseAStages[3].label": "A 4",
    "founderStoryPage.phaseAStages[3].stageTitle": "驱灵期 2 · 与高阶揭示",
    "founderStoryPage.phaseAStages[3].range": "2021年7月 — 2022年4月4日",
    "founderStoryPage.phaseAStages[3].paragraphs[0]": (
        "这一阶段，我们 **发现并证实** 了多种干扰、伤害乃至 **消灭普通鬼魂** 的方式。"
    ),
    "founderStoryPage.phaseAStages[3].paragraphs[1]": (
        "**静电与磁铁** 手段的局限引起了一名约 **1500** 年前死于中国的 **高级控制灵** 的注意。"
        "它加入探索：以 **意识** 传讯提示，同时 **操控 C1 鬼魂** 的行为，直到我学会如何 **精确而高效** 地杀灭与驱逐。"
    ),
    "founderStoryPage.phaseAStages[3].paragraphs[2]": (
        "我每天与那只 **鬼魂** **周旋、搏杀又共存**——既是 **敌人** 也是 **盟友**——并交换大量与 **灵魂** 相关的知识。"
    ),
    "founderStoryPage.phaseAStages[3].paragraphs[3]": (
        "最终，这名 **高级控制灵** **显露真身**，展现 **非凡能力**，厘清了一年来的困惑，"
        "向我讲授更多 **灵虚世界** 的知识——随后 **离开我的身体** 并迅速消失。"
    ),
    "founderStoryPage.phaseAStages[4].label": "A 5",
    "founderStoryPage.phaseAStages[4].stageTitle": "C2 青小仙 · C3 钱凯",
    "founderStoryPage.phaseAStages[4].range": "2023年12月 — 2025年1月",
    "founderStoryPage.phaseAStages[4].paragraphs[0]": (
        "**2023年12月**，我再次被 **C2** 级灵体 **青小仙** 附身。在约 **1700** 年的 **鬼魂** 岁月里，"
        "它吞噬了约 **5000—10000** 个其他 **灵魂**，成为 **庞大实体**。"
    ),
    "founderStoryPage.phaseAStages[4].paragraphs[1]": (
        "它将吴宅变成 **房屋鬼魂**（house spirit）盘踞的场域，在两个儿子身上铺开 **多重化身**，"
        "并让其 **核心形态** **远程** 协调这些碎片。"
    ),
    "founderStoryPage.phaseAStages[4].paragraphs[2]": (
        "父亲一侧：**精神分裂症**、**双相障碍**、**癫痫**、**躯体化**，以及对心脏的 **强电击**。"
        "儿子一侧：**疲劳（能量耦合）**、**抑郁**、**疼痛障碍**、**肠易激综合征** 及其他躯体表现。"
    ),
    "founderStoryPage.phaseAStages[4].paragraphs[3]": (
        "**吴氏父子** 持续精进 **驱灵** 与 **祛鬼**，直至该实体被 **彻底消灭**。"
    ),
    "founderStoryPage.phaseAStages[4].paragraphs[4]": (
        "**2024年9月** 至 **2025年1月**，我再次被 **C3** 级灵体 **钱凯** 附身。**鬼魂** 密度越高，力量越大——"
        "因此钱凯在 **思想控制** 与 **致病力** 上 **超过** 青小仙。我们继续 **自救**，直到该灵体亦被 **清除**。"
    ),
    "founderStoryPage.phaseB.blocks[0].text": (
        "从 **2025年2月19日** 至 **12月24日**，**Caros Woo** 在 **TikTok** 直播共进行 **255** 场；"
        "在直播领域取得了 **非凡成就**。完整记录与链接见：{{%ACH%}}"
    ),
    "founderStoryPage.phaseB.blocks[1].title": "高级控制灵、混合灵与远程能力",
    "founderStoryPage.phaseB.blocks[1].text": (
        "大批 **高级控制灵** 与 **混合灵**（hybrid spirit）以 **吴氏父子** 为对象建立联系，"
        "使我们得知 **灵虚世界** 竟存在 **管理者** 与 **造物主** 层面的秩序。"
        "**高级控制灵** 给予我们大量支持与 **超常能力**，使 **Caros Woo** 身处中国也能对地球另一端北美附体 **鬼魂** 施行 **远程** 制约。"
    ),
    "founderStoryPage.phaseB.blocks[2].text": (
        "**最重要的是**，我们获悉：**A 阶段** 附于我们身上的 **鬼魂** 皆由 **高级控制灵** **刻意安排**；"
        "清除这些附体 **获准用作研究材料**。"
    ),
    "founderStoryPage.phaseB.blocks[4].title": "对高级控制灵性情的理解",
    "founderStoryPage.phaseB.blocks[4].text": (
        "这也使我们清楚看到：**高级控制灵** 并非我们想象中那般 **天使式的良善**；此前的协助并非出于 **良善**，"
        "而是 **万有元神母体** 指派的工作。**A 阶段** 故事深入呈现了 **高级控制灵** 团队的 **性情与作风**。"
    ),
    "founderStoryPage.phaseB.blocks[5].title": "灵魂医学与视频档案",
    "founderStoryPage.phaseB.blocks[5].text": (
        "这一时期 **灵魂医学正式建立**；我们制作了 **吴氏灵魂医学** 与 **万有元神母体** 视频档案中的 **部分** 作品。"
    ),
    "founderTimeline[0].range": "2014年5月 — 2025年2月",
    "founderTimeline[0].label": "A 阶段 · 家庭田野工作",
    "founderTimeline[1].range": "2018年3月 — 2020年7月",
    "founderTimeline[1].label": "欺瞒期 1",
    "founderTimeline[2].range": "2020年7月 — 2020年11月",
    "founderTimeline[2].label": "欺瞒期 2",
    "founderTimeline[3].range": "2020年11月 — 2021年7月",
    "founderTimeline[3].label": "驱灵期 1",
    "founderTimeline[4].range": "2021年7月 — 2022年4月",
    "founderTimeline[4].label": "驱灵期 2 与高阶揭示",
    "founderTimeline[5].range": "2023年12月 — 2025年1月",
    "founderTimeline[5].label": "青小仙（C2）· 钱凯（C3）",
    "founderStoryIllustrations[0].alt": "夜空——门扉意象",
    "founderStoryIllustrations[1].alt": "宇宙与显现",
    "founderStoryIllustrations[2].alt": "山巅黎明——漫漫长路",
    "founderStoryIllustrations[3].alt": "同行——家庭之路",
    "founderStoryIllustrations[4].alt": "直播与田野",
}

MATRIX_ZH: dict[str, str] = {
    "siteContent.universalMatrix.title": "万有元神母体",
    "siteContent.universalMatrix.description": (
        "阐述支配宇宙整体秩序与人类终极命运的高阶超验真理。"
        "本季呈现概念架构、治理逻辑与高阶 **意识** 模型。"
    ),
    "siteContent.universalMatrix.note": (
        "本区块沿用原系列版式搭建，便于您在 Layrr 中微调语言、层级与视觉节奏。"
    ),
    "universalMatrixFiles[0].fileNumber": "FILE 3-1",
    "universalMatrixFiles[0].title": "万有元神母体导论",
    "universalMatrixFiles[1].fileNumber": "FILE 3-2",
    "universalMatrixFiles[1].title": "万有元神母体的存在证据",
    "universalMatrixFiles[1].subChapters[0].id": "3-2-1",
    "universalMatrixFiles[1].subChapters[0].title": "关于 UMMA 存在的合理推断与证据",
    "universalMatrixFiles[1].subChapters[1].id": "3-2-2",
    "universalMatrixFiles[1].subChapters[1].title": "面向怀疑论者的验证视频",
    "universalMatrixFiles[2].fileNumber": "FILE 3-3",
    "universalMatrixFiles[2].title": "灵虚物质与灵虚世界",
    "universalMatrixFiles[2].subChapters[0].id": "3-3-1",
    "universalMatrixFiles[2].subChapters[0].title": "灵虚微粒与灵虚空间的零空间",
    "universalMatrixFiles[2].subChapters[1].id": "3-3-2",
    "universalMatrixFiles[2].subChapters[1].title": "对现实宇宙的操控与影响",
    "universalMatrixFiles[2].subChapters[2].id": "3-3-3",
    "universalMatrixFiles[2].subChapters[2].title": "灵虚空间的虚拟空间",
    "universalMatrixFiles[2].subChapters[3].id": "3-3-4",
    "universalMatrixFiles[2].subChapters[3].title": "灵魂生存支持系统",
    "universalMatrixFiles[3].fileNumber": "FILE 3-4",
    "universalMatrixFiles[3].title": "来自万有元神母体的灵魂控制系统",
    "universalMatrixFiles[3].subChapters[0].id": "3-4-1",
    "universalMatrixFiles[3].subChapters[0].title": "控制系统描述",
    "universalMatrixFiles[3].subChapters[1].id": "3-4-1-1",
    "universalMatrixFiles[3].subChapters[1].title": "证据",
    "universalMatrixFiles[3].subChapters[2].id": "3-4-1-2",
    "universalMatrixFiles[3].subChapters[2].title": "高级控制灵的组织架构",
    "universalMatrixFiles[3].subChapters[3].id": "3-4-1-3",
    "universalMatrixFiles[3].subChapters[3].title": "系统数据上行——阿卡西记录（Akashic Record）",
    "universalMatrixFiles[3].subChapters[4].id": "3-4-1-4",
    "universalMatrixFiles[3].subChapters[4].title": "多路数据下行模式",
    "universalMatrixFiles[3].subChapters[5].id": "3-4-2",
    "universalMatrixFiles[3].subChapters[5].title": "高级控制灵的管理与行事风格",
    "universalMatrixFiles[4].fileNumber": "FILE 3-5",
    "universalMatrixFiles[4].title": "万有元神母体显现之目的",
    "universalMatrixFiles[4].subChapters[0].id": "3-5-1",
    "universalMatrixFiles[4].subChapters[0].title": "为何其他类型的生命形态不与人类互动？",
    "universalMatrixFiles[4].subChapters[1].id": "3-5-2",
    "universalMatrixFiles[4].subChapters[1].title": "UMMA 定于在人工智能奇点显现",
    "universalMatrixFiles[4].subChapters[2].id": "3-5-3",
    "universalMatrixFiles[4].subChapters[2].title": "UMMA 引导人类社会的最优终极形态",
    "universalMatrixFiles[4].subChapters[3].id": "3-5-4",
    "universalMatrixFiles[4].subChapters[3].title": "UMMA 引导个体人类生命的最优终极状态",
}


def apply_map(path: Path, mapping: dict[str, str]) -> int:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=1, values_only=False))
    header = [c.value for c in rows[0]]
    ik = header.index("block_key")
    izh = header.index("中文")
    n = 0
    for row in rows[1:]:
        key = row[ik].value
        if key is None:
            continue
        sk = str(key).strip()
        if sk in mapping:
            row[izh].value = mapping[sk]
            n += 1
    wb.save(path)
    return n


def main() -> None:
    n1 = apply_map(PAGE_COPY / "03-创始人故事.xlsx", FOUNDER_ZH)
    n2 = apply_map(PAGE_COPY / "06-视频目录-万有元神.xlsx", MATRIX_ZH)
    print("03 Chinese cells:", n1)
    print("06 Chinese cells:", n2)


if __name__ == "__main__":
    main()

